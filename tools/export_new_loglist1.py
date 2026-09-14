import csv
from openpyxl import load_workbook

# =========================================================
# CONFIG
# =========================================================

XLSX = "INPUT_ANGKUTAN_STOCK_NEW.xlsx"
SHEET = "POSISI TERAKHIR"

# Kolom posisi terakhir ada di T
POSITION_COL = 20  # T

# Output AH sampai AQ
OUT_MIN_COL = 34   # AH
OUT_MAX_COL = 43   # AQ

# Supaya bisa membaca posisi T sekaligus data AH:AQ
READ_MIN_COL = 20  # T
READ_MAX_COL = 43  # AQ

# Baris 3 adalah header
MIN_ROW = 3

# Output tetap loglist1.csv
OUT_CSV = "loglist1.csv"


# =========================================================
# HELPER
# =========================================================

def cell_str(value):
    """
    Merapikan nilai cell sebelum ditulis ke CSV.
    """
    if value is None:
        return ""

    if isinstance(value, str):
        value = value.strip()

        if value == "=":
            return ""

        return value

    return str(value)


def is_invalid_nobtg(value):
    """
    Abaikan baris jika noBtg kosong atau 0.
    """
    if value is None:
        return True

    if isinstance(value, (int, float)):
        return value == 0

    value = str(value).strip()

    return value in ("", "0", "0.0")


def should_skip_posisi(value):
    """
    Skip data jika posisi terakhir:
    - DKDS
    - mengandung kata MILIR
    """
    if value is None:
        return False

    posisi = str(value).strip().upper()

    if posisi == "DKDS":
        return True

    if "MILIR" in posisi:
        return True

    return False


# =========================================================
# MAIN
# =========================================================

def main():

    print("=" * 60)
    print("NEW LOGLIST1 EXPORTER")
    print("=" * 60)

    print(f"Input Excel : {XLSX}")
    print(f"Sheet       : {SHEET}")
    print(f"Output      : {OUT_CSV}")
    print()

    # -----------------------------------------------------
    # Buka Excel
    # -----------------------------------------------------

    print("Membuka Excel...")

    wb = load_workbook(
        XLSX,
        read_only=True,
        data_only=True
    )

    if SHEET not in wb.sheetnames:

        wb.close()

        raise SystemExit(
            f"ERROR: Sheet '{SHEET}' tidak ditemukan.\n"
            f"Sheet tersedia: {wb.sheetnames}"
        )

    ws = wb[SHEET]

    print(f"Baris terakhir sheet terdeteksi : {ws.max_row}")
    print(f"Kolom terakhir sheet terdeteksi : {ws.max_column}")
    print()

    # -----------------------------------------------------
    # Posisi AH:AQ di dalam range T:AQ
    # -----------------------------------------------------

    out_start = OUT_MIN_COL - READ_MIN_COL

    out_end = (
        out_start
        + (OUT_MAX_COL - OUT_MIN_COL + 1)
    )

    # -----------------------------------------------------
    # Counter
    # -----------------------------------------------------

    total_excel = 0
    total_output = 0

    skip_no_btg = 0
    skip_dkds = 0
    skip_milir = 0

    # -----------------------------------------------------
    # Export
    # -----------------------------------------------------

    with open(
        OUT_CSV,
        "w",
        newline="",
        encoding="utf-8"
    ) as file:

        writer = csv.writer(file)

        for row_number, row in enumerate(
            ws.iter_rows(
                min_row=MIN_ROW,
                min_col=READ_MIN_COL,
                max_col=READ_MAX_COL,
                values_only=True
            ),
            start=MIN_ROW
        ):

            # Kolom T
            posisi_raw = row[0]

            # AH sampai AQ
            out_row = row[out_start:out_end]

            # =================================================
            # HEADER
            # =================================================

            if row_number == MIN_ROW:

                writer.writerow(
                    [cell_str(v) for v in out_row]
                )

                continue

            total_excel += 1

            # AH = noBtg
            nobtg_raw = out_row[0]

            # =================================================
            # FILTER noBtg
            # =================================================

            if is_invalid_nobtg(nobtg_raw):

                skip_no_btg += 1

                continue

            # =================================================
            # FILTER POSISI
            # =================================================

            if posisi_raw is not None:

                posisi = str(
                    posisi_raw
                ).strip().upper()

                if posisi == "DKDS":

                    skip_dkds += 1

                    continue

                if "MILIR" in posisi:

                    skip_milir += 1

                    continue

            # =================================================
            # WRITE
            # =================================================

            writer.writerow(
                [cell_str(v) for v in out_row]
            )

            total_output += 1

    wb.close()

    # -----------------------------------------------------
    # REPORT
    # -----------------------------------------------------

    print()
    print("=" * 60)
    print("HASIL EXPORT")
    print("=" * 60)

    print(f"Baris diperiksa       : {total_excel}")
    print(f"Skip No.Btg kosong/0  : {skip_no_btg}")
    print(f"Skip DKDS             : {skip_dkds}")
    print(f"Skip MILIR            : {skip_milir}")
    print(f"Data masuk CSV        : {total_output}")

    print()
    print(f"Export selesai -> {OUT_CSV}")


if __name__ == "__main__":
    main()
