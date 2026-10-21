import csv, os, sys, json
from datetime import timezone
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook

SOURCE = os.getenv('SOURCE_XLSX', 'data ukur LHP IBS.xlsx')
OUTPUT = 'stock_olahan_ibs.csv'

def numeric(v, cell):
    if v is None or v == '':
        raise ValueError(f'{cell} kosong: pastikan file Excel telah dihitung dan disimpan.')
    try:
        return Decimal(str(v).replace(',', '.'))
    except InvalidOperation:
        raise ValueError(f'{cell} bukan angka: {v!r}. Pastikan cache rumus Excel tersedia.')

def fmt(v, places=4):
    return f'{v:.{places}f}'

def main():
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    if 'REKAP' not in wb.sheetnames:
        raise ValueError('Sheet REKAP tidak ditemukan')
    ws = wb['REKAP']
    stock = [numeric(ws[c].value, c) for c in ('K7', 'L7', 'M7')]
    extra = [numeric(ws[c].value,c) for c in ('C32','D32','D33','D34')]
    details = []
    for cells in ws.iter_rows(min_row=5, max_row=252, min_col=15, max_col=20, values_only=True):
        no, tanggal, dko, seri, vol, tujuan = cells
        if not any(x is not None and str(x).strip() for x in cells):
            continue
        # Ignore template rows with a number but no actual transaction.
        if not dko and not seri and not vol and not tujuan:
            continue
        if not dko or not seri or vol is None:
            raise ValueError(f'Detail DKO belum lengkap: nomor urut {no!r}')
        if isinstance(tanggal, (datetime, date)):
            tanggal = tanggal.strftime('%d-%m-%Y')
        else:
            tanggal = str(tanggal or '').strip()
        details.append(['DETAIL', '', '', '', str(no or ''), tanggal, str(dko).strip(), str(seri).strip(), fmt(numeric(vol, f'Volume DKO {no}')), str(tujuan or '').strip()])
    with open(OUTPUT + '.tmp', 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        writer.writerow(['tipe','stock_awal_m3','pengurangan_m3','stock_akhir_m3','no','tanggal','no_dko','seri_skshhk','volume_m3','tujuan','belum_proses_batang','belum_proses_m3','rencana_sortiment_m3','total_stock_olahan_m3'])
        writer.writerow(['RINGKASAN', *[fmt(x) for x in stock], '', '', '', '', '', '', *[fmt(x) for x in extra]])
        writer.writerows([r+['']*4 for r in details])
    os.replace(OUTPUT + '.tmp', OUTPUT)
    # Optional PHAT sheet: export exact N1:AK54 as JSON, preserving 54 x 24 grid.
    phat_file='monitoring_phat_ibs.json'
    if 'RKP-LHP' in wb.sheetnames:
        phat=wb['RKP-LHP']
        def serial(v):
            if isinstance(v,(datetime,date)):return v.isoformat()
            if isinstance(v,Decimal):return float(v)
            return v
        rows=[[serial(v) for v in row] for row in phat.iter_rows(min_row=1,max_row=54,min_col=14,max_col=37,values_only=True)]
        with open(phat_file+'.tmp','w',encoding='utf-8') as f:
            json.dump({'source':'RKP-LHP!N1:AK54','updated_at':datetime.now().astimezone().isoformat(timespec='seconds'),'rows':rows},f,ensure_ascii=False,default=str)
        os.replace(phat_file+'.tmp',phat_file)
        print('Monitoring PHAT exported:',phat_file)
    else:
        # Remove stale PHAT data when source sheet is absent.
        if os.path.exists(phat_file):os.remove(phat_file)
        print('INFO: Sheet RKP-LHP tidak ada; Monitoring PHAT menunggu sheet sumber.')
    print(f'{OUTPUT}: stok {list(map(fmt,stock))}; {len(details)} DKO')

if __name__ == '__main__':
    try: main()
    except Exception as e: print('ERROR:', e, file=sys.stderr); sys.exit(1)
